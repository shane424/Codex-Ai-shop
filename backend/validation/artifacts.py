from dataclasses import dataclass
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

@dataclass
class ValidationResult:
    passed: bool; reasons: list[str]

def validate_xlsx(path: Path, required=('Instructions','Data Entry','Dashboard')) -> ValidationResult:
    reasons=[]
    try:
        wb=load_workbook(path,data_only=False)
        reasons += [f'missing sheet: {x}' for x in required if x not in wb.sheetnames]
        if 'Data Entry' in wb:
            ws=wb['Data Entry'];
            if ws.max_row<2 or ws.max_column<2: reasons.append('insufficient dimensions')
            if any(c.value in (None,'') for c in ws[1]): reasons.append('blank header')
        if 'Dashboard' in wb and not any(isinstance(c.value,str) and c.value.startswith('=') for row in wb['Dashboard'] for c in row): reasons.append('no formulas')
    except Exception as exc: reasons.append(f'unreadable workbook: {type(exc).__name__}')
    return ValidationResult(not reasons,reasons)
def validate_pdf(path: Path) -> ValidationResult:
    reasons=[]
    try:
        reader=PdfReader(path)
        if not reader.pages: reasons.append('no pages')
        for i,page in enumerate(reader.pages):
            if not (page.extract_text() or '').strip(): reasons.append(f'blank page: {i+1}')
    except Exception as exc: reasons.append(f'unreadable PDF: {type(exc).__name__}')
    return ValidationResult(not reasons,reasons)
