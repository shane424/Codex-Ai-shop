import re
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen.canvas import Canvas


@dataclass(frozen=True)
class ProductSpec:
    name: str
    niche: str
    kind: str = "xlsx"
    price_cents: int = 1200

    @property
    def slug(self) -> str:
        return re.sub(r"[^a-z0-9]+", "-", self.name.lower()).strip("-")


CATALOG = [
    ProductSpec("Contractor Job Profit Tracker", "Small Business", price_cents=1900),
    ProductSpec("Small Business Expense Tracker", "Small Business"),
    ProductSpec("Equipment Maintenance Log", "Small Business", "pdf", 500),
    ProductSpec("Customer Follow-Up CRM", "Small Business"),
    ProductSpec("Quote Comparison Spreadsheet", "Small Business"),
    ProductSpec("Invoice Payment Tracker", "Small Business"),
    ProductSpec("Inventory Reorder Calculator", "Small Business", price_cents=1900),
    ProductSpec("Mileage Log", "Small Business", price_cents=900),
    ProductSpec("Vendor Comparison Sheet", "Small Business"),
    ProductSpec("Simple Project Profitability Calculator", "Small Business", price_cents=1900),
    ProductSpec("Pet Medication Tracker", "Pet", price_cents=900),
    ProductSpec("Dog Daycare Daily Report", "Pet", "pdf", 500),
    ProductSpec("Kennel Cleaning Checklist", "Pet", "pdf", 500),
    ProductSpec("Pet Sitter Client Intake Pack", "Pet", "pdf", 900),
    ProductSpec("Reptile Feeding Tracker", "Pet", price_cents=900),
    ProductSpec("Job Application Tracker", "Job & Freelance", price_cents=900),
    ProductSpec("Recruiter CRM", "Job & Freelance"),
    ProductSpec("Freelance Lead Tracker", "Job & Freelance"),
    ProductSpec("Client Profitability Calculator", "Job & Freelance", price_cents=1900),
    ProductSpec("Proposal Follow-Up Tracker", "Job & Freelance"),
]

BUNDLES = [
    ("Small Business Operations Pack", 3900, range(1, 11)),
    ("Contractor Starter Pack", 2900, [1, 3, 5, 8]),
    ("Pet Business Operations Pack", 2900, range(11, 16)),
    ("Job Search Command Center", 1900, [16, 17]),
    ("Freelancer Operations Pack", 2900, [18, 19, 20]),
]


def opportunity_score(utility: float, purchase_intent: float, automation: float, margin: float, specificity: float) -> float:
    values = (utility, purchase_intent, automation, margin, specificity)
    if any(not 0 <= value <= 100 for value in values):
        raise ValueError("Scores must be between 0 and 100")
    return round(utility * .30 + purchase_intent * .25 + automation * .20 + margin * .15 + specificity * .10, 2)


def generate_xlsx(spec: ProductSpec, destination: Path) -> None:
    wb = Workbook()
    instructions = wb.active
    instructions.title = "Instructions"
    instructions.append([spec.name])
    instructions.append(["Enter one item per row on the Tracker sheet. Summary formulas update automatically."])
    tracker = wb.create_sheet("Tracker")
    headers = ["Date", "Item", "Category", "Quantity", "Unit Amount", "Total", "Status", "Notes"]
    tracker.append(headers)
    for row in range(2, 102):
        tracker.cell(row, 6, f'=IFERROR(D{row}*E{row},0)')
    tracker.freeze_panes = "A2"
    tracker.auto_filter.ref = "A1:H101"
    tracker.column_dimensions["B"].width = 28
    tracker.column_dimensions["H"].width = 36
    for cell in tracker[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245B78")
    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Entries", '=COUNTA(Tracker!B2:B101)'])
    summary.append(["Total", '=SUM(Tracker!F2:F101)'])
    summary.append(["Completed", '=COUNTIF(Tracker!G2:G101,"Complete")'])
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)


def generate_pdf(spec: ProductSpec, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(str(destination), pagesize=letter)
    canvas.setTitle(spec.name)
    canvas.setFont("Helvetica-Bold", 18)
    canvas.drawString(54, 744, spec.name)
    canvas.setFont("Helvetica", 10)
    canvas.drawString(54, 722, "Reusable practical worksheet — print one copy for each record.")
    y = 680
    for label in ["Date", "Name / item", "Owner", "Status", "Details", "Action required", "Follow-up date", "Notes"]:
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawString(54, y, label)
        canvas.line(160, y - 2, 550, y - 2)
        y -= 55
    canvas.setFont("Helvetica", 8)
    canvas.drawString(54, 36, "Utility Shop • General record-keeping template • Not professional advice")
    canvas.save()


def validate_product(path: Path) -> tuple[bool, str]:
    if not path.exists() or path.stat().st_size < 100:
        return False, "Missing or empty artifact"
    try:
        if path.suffix == ".xlsx":
            wb = load_workbook(path, data_only=False)
            if not {"Instructions", "Tracker", "Summary"}.issubset(wb.sheetnames):
                return False, "Required worksheets missing"
            if not any(cell.data_type == "f" for row in wb["Tracker"] for cell in row):
                return False, "No formulas found"
        elif path.suffix == ".pdf":
            reader = PdfReader(str(path))
            if not reader.pages or not "".join(page.extract_text() or "" for page in reader.pages).strip():
                return False, "PDF has no readable content"
        elif path.suffix == ".zip":
            with ZipFile(path) as archive:
                if not archive.namelist():
                    return False, "Bundle is empty"
        else:
            return False, "Unsupported file type"
    except Exception as exc:
        return False, f"Unreadable artifact: {exc}"
    return True, "ok"


def generate_artifact(spec: ProductSpec, directory: Path) -> Path:
    path = directory / f"{spec.slug}.{spec.kind}"
    if not path.exists():
        (generate_pdf if spec.kind == "pdf" else generate_xlsx)(spec, path)
    valid, reason = validate_product(path)
    if not valid:
        raise ValueError(f"QA failed for {spec.name}: {reason}")
    return path


def generate_bundle(name: str, indexes, directory: Path) -> Path:
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    path = directory / f"{slug}.zip"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for index in indexes:
            spec = CATALOG[index - 1]
            artifact = generate_artifact(spec, directory)
            archive.write(artifact, artifact.name)
    return path

