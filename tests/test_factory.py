from zipfile import ZipFile
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest
from backend.factory import ProductSpec, generate_artifact, generate_bundle, opportunity_score, validate_product, CATALOG


def test_opportunity_score_and_validation():
    assert opportunity_score(80, 80, 80, 80, 80) == 80
    with pytest.raises(ValueError):
        opportunity_score(101, 80, 80, 80, 80)


def test_spreadsheet_generation_and_qa(tmp_path):
    path = generate_artifact(ProductSpec("Test Tracker", "Test"), tmp_path)
    assert validate_product(path) == (True, "ok")
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["Instructions", "Tracker", "Summary"]
    assert workbook["Tracker"]["F2"].data_type == "f"


def test_pdf_generation_and_qa(tmp_path):
    path = generate_artifact(ProductSpec("Test Checklist", "Test", "pdf", 500), tmp_path)
    assert validate_product(path) == (True, "ok")
    assert len(PdfReader(path).pages) == 1


def test_bundle_contains_expected_safe_files(tmp_path):
    path = generate_bundle("Test Bundle", [1, 2], tmp_path)
    with ZipFile(path) as archive:
        names = archive.namelist()
    assert len(names) == 2
    assert all("/" not in name and ".." not in name for name in names)

